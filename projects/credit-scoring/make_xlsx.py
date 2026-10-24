import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

S = json.load(open('/home/claude/portfolio/projects/banking-sector/sector.json'))
P = 'Q4 2024'; D = S['data'][P]; NAMES = S['names']
codes = sorted(D, key=lambda c: -D[c]['ta'])

F = 'Arial'
def font(color='000000', b=False, sz=10, i=False): return Font(name=F, color=color, bold=b, size=sz, italic=i)
BLUE, GREEN, BLACK = '0000FF', '008000', '000000'
YEL = PatternFill('solid', fgColor='FFFF00'); HDR = PatternFill('solid', fgColor='0F2A47'); SUB = PatternFill('solid', fgColor='E8EDF3'); OUT = PatternFill('solid', fgColor='EAF2FB')
thin = Side(style='thin', color='C3C2B7'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); BOT = Border(bottom=thin)
MN = '#,##0;(#,##0);-'; PC = '0.0%;(0.0%);-'; PC2 = '0.00%;(0.00%);-'

wb = Workbook()
def sheet(name, first=False):
    ws = wb.active if first else wb.create_sheet(name); ws.title = name; ws.sheet_view.showGridLines = False
    return ws
def hdr(ws, r, vals, c0=1):
    for j, v in enumerate(vals):
        c = ws.cell(r, c0 + j, v); c.font = font('FFFFFF', True); c.fill = HDR; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOX
def put(ws, ref, v, color=BLACK, b=False, fmt=None, fill=None, i=False, sz=10, wrap=False, align=None):
    c = ws[ref]; c.value = v; c.font = font(color, b, sz, i)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if wrap or align: c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical='top' if wrap else None)
    return c

# ---------------- Cover ----------------
ws = sheet('Cover', True)
ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 26; ws.column_dimensions['C'].width = 90
put(ws, 'B2', 'Bank Counterparty Credit Scoring Model', BLACK, True, sz=16)
put(ws, 'B3', 'Author: Jamal Kanbarieh  ·  Built for portfolio demonstration using public disclosures', '52514E', i=True)
rows = [
 ('Purpose', 'Assigns an internal credit grade and an indicative probability of default (PD) to a bank counterparty, as used by a financial-institutions desk to set interbank and trade-finance limits.'),
 ('Approach', 'CAMELS-style scorecard: 9 ratios across Capital, Asset quality, Profitability, Liquidity and Efficiency. Each ratio is graded A (1 point, strongest) to E (5 points) against thresholds; category weights give a score from 1 to 5, which maps to a rating scale and PD. An optional sovereign ceiling caps the final rating.'),
 ('How to use', '1) On "Inputs", choose a bank in the yellow cell C4 (data is pulled from "Peer Data"). 2) Optionally type over any figure in the yellow override column. 3) Read the result on "Scorecard". 4) "Peer Ranking" scores every bank at once. 5) Thresholds, weights and the rating scale are editable on "Assumptions".'),
 ('Colour legend', 'Blue text = hard-coded input · Black = formula · Green = link to another sheet · Yellow fill = cell for you to choose/override.'),
 ('Data source', f'Public quarterly financial statements of Syrian private banks ({P}), published by the Syrian Commission on Financial Markets and Securities (scfms.sy/en/declaration). Figures in SYP millions. Net income and operating income exclude unrealised structural FX revaluation.'),
 ('Calibration', 'Thresholds, weights and the score-to-rating-to-PD mapping are the author\'s own illustrative calibration. They are not an official methodology of any bank or rating agency, and the PDs are not historical default statistics.'),
 ('Disclaimer', 'Demonstration model only. Not a credit opinion or investment advice.'),
]
for k, (a, b) in enumerate(rows):
    r = 5 + k * 2; put(ws, f'B{r}', a, BLACK, True); put(ws, f'C{r}', b, wrap=True); ws.row_dimensions[r].height = 46
put(ws, 'B20', 'Example input format', BLACK, True); put(ws, 'C20', 'Inputs!C4 = "QNBS" (bank code from the list) · Inputs!C7 = 12 (months covered by the income statement) · overrides in SYP millions, e.g. 3,517,320', '52514E', i=True, wrap=True)

# ---------------- Peer Data ----------------
pdw = sheet('Peer Data')
FIELDS = [('Total assets', 'ta'), ('Liquid assets', 'liq'), ('Gross loans (direct facilities)', 'gl'), ('Non-performing loans', 'npl'), ('ECL provisions on loans', 'ecl'),
          ('Customer deposits', None), ('Cash margins', 'cm'), ('Total liabilities', 'tl'), ('Total equity', 'eqi'), ('Regulatory capital', 'regcap'), ('Risk-weighted assets', 'rwa'),
          ('Operating income (adj.)', 'goi'), ('Operating expenses (excl. ECL)', None), ('Net income (adj.)', 'ni')]
put(pdw, 'A1', f'Peer data: Syrian private banks, {P} (SYP millions)', BLACK, True, sz=12)
put(pdw, 'A2', 'Source: SCFMS public quarterly disclosures (scfms.sy/en/declaration), compiled by the author. Year-to-date income = 12 months. Banks that had not published by the cut-off (CHB, SIIB, SGB) are excluded.', '52514E', i=True)
hdr(pdw, 4, ['Code', 'Bank', 'Period'] + [f[0] for f in FIELDS])
for i, c in enumerate(codes):
    r = 5 + i; d = D[c]
    put(pdw, f'A{r}', c, BLUE); put(pdw, f'B{r}', NAMES[c], BLUE); put(pdw, f'C{r}', P, BLUE)
    for j, (lab, k) in enumerate(FIELDS):
        v = (d['ca'] + d['td']) if lab == 'Customer deposits' else abs(d['opx']) if lab.startswith('Operating expenses') else d[k]
        put(pdw, f'{L(4 + j)}{r}', round(v / 1e6, 1), BLUE, fmt=MN)
NB = len(codes); LAST = 4 + NB
pdw.column_dimensions['A'].width = 9; pdw.column_dimensions['B'].width = 30; pdw.column_dimensions['C'].width = 9
for j in range(len(FIELDS)): pdw.column_dimensions[L(4 + j)].width = 15
pdw.row_dimensions[4].height = 42; pdw.freeze_panes = 'D5'

# ---------------- Assumptions ----------------
asw = sheet('Assumptions')
put(asw, 'A1', 'Scoring assumptions (editable)', BLACK, True, sz=12)
put(asw, 'A2', "Author's illustrative calibration. A = 1 point (strongest) … E = 5 points. A ratio is graded A if it meets the A cut-off, B if it meets the B cut-off, and so on; otherwise E.", '52514E', i=True)
hdr(asw, 4, ['Ratio', 'Category', 'Better when', 'A cut-off', 'B cut-off', 'C cut-off', 'D cut-off', 'Weight', 'Definition'])
RAT = [
 ('Capital adequacy ratio (CAR)', 'Capital', 'Higher', .20, .15, .125, .10, .10, 'Regulatory capital ÷ risk-weighted assets'),
 ('Equity / RWA', 'Capital', 'Higher', .15, .12, .10, .08, .10, 'Total equity ÷ risk-weighted assets (loss-absorption buffer)'),
 ('NPL ratio', 'Asset quality', 'Lower', .03, .06, .10, .20, .15, 'Non-performing loans ÷ gross loans'),
 ('NPL / (equity + ECL)', 'Asset quality', 'Lower', .10, .25, .50, 1.00, .15, 'Unprovisioned-risk test ("Texas ratio" style); n.m. scores E'),
 ('Net income / RWA (annualised)', 'Profitability', 'Higher', .025, .015, .0075, 0, .10, 'Adjusted net income × 12 / months ÷ RWA'),
 ('Pre-provision profit / RWA (ann.)', 'Profitability', 'Higher', .04, .025, .015, 0, .10, '(Operating income − operating expenses) × 12 / months ÷ RWA'),
 ('Market-funding gap / assets', 'Liquidity', 'Lower', -.20, -.10, 0, .10, .10, '(Liabilities − customer deposits − cash margins − liquid assets) ÷ total assets'),
 ('Liquid assets / customer deposits', 'Liquidity', 'Higher', .60, .45, .30, .20, .10, 'Liquid assets ÷ customer deposits'),
 ('Cost-to-income ratio', 'Efficiency', 'Lower', .45, .55, .65, .80, .10, 'Operating expenses ÷ operating income; n.m. scores E'),
]
for i, rr in enumerate(RAT):
    r = 5 + i
    for j, v in enumerate(rr):
        col = L(1 + j); fmt = PC if j in (3, 4, 5, 6, 7) else None
        put(asw, f'{col}{r}', v, BLUE if j >= 2 and j <= 7 else BLACK, fmt=fmt)
dv = DataValidation(type='list', formula1='"Higher,Lower"', allow_blank=False); asw.add_data_validation(dv); dv.add('C5:C13')
put(asw, 'G14', 'Total weight', BLACK, True); put(asw, 'H14', '=SUM(H5:H13)', BLACK, True, fmt=PC)
put(asw, 'I14', '=IF(ABS(H14-1)<0.0001,"OK: weights sum to 100%","CHECK: weights must sum to 100%")', BLACK, i=True)

put(asw, 'A17', 'Rating scale (score → rating → indicative PD)', BLACK, True, sz=11)
hdr(asw, 18, ['Score from', 'Rating', 'Indicative 1-yr PD'])
SCALE = [(1.0, 'AA', .0003), (1.4, 'A', .0006), (1.8, 'BBB+', .0010), (2.1, 'BBB', .0017), (2.4, 'BBB-', .0027), (2.7, 'BB+', .0045), (3.0, 'BB', .0075), (3.3, 'BB-', .013), (3.6, 'B+', .022), (3.9, 'B', .038), (4.2, 'B-', .065), (4.5, 'CCC', .25)]
for i, (a, b, c) in enumerate(SCALE):
    r = 19 + i; put(asw, f'A{r}', a, BLUE, fmt='0.00'); put(asw, f'B{r}', b, BLUE); put(asw, f'C{r}', c, BLUE, fmt=PC2)
asw['C19'].comment = Comment("Illustrative PD calibration by the author. It is monotonic and broadly consistent in ordering with published long-run agency default studies, but it is not an official statistic.", 'Jamal Kanbarieh')
SC0, SC1 = 19, 18 + len(SCALE)
put(asw, 'E18', 'Sovereign ceiling options', BLACK, True); put(asw, f'E19', 'None', BLUE)
for i, (_, b, _) in enumerate(SCALE): put(asw, f'E{20 + i}', b, BLUE)
for col, w in zip('ABCDEFGHI', [34, 15, 12, 11, 11, 11, 11, 9, 70]): asw.column_dimensions[col].width = w
asw.row_dimensions[4].height = 30

# ---------------- Inputs ----------------
inw = sheet('Inputs')
put(inw, 'A1', 'Inputs', BLACK, True, sz=12)
put(inw, 'A2', 'Choose a bank in C4. Figures are pulled from "Peer Data"; type a value in the yellow Override column to replace any of them. SYP millions.', '52514E', i=True)
put(inw, 'B4', 'Bank code', BLACK, True); put(inw, 'C4', 'QNBS', BLUE, True, fill=YEL)
dv2 = DataValidation(type='list', formula1=f"='Peer Data'!$A$5:$A${LAST}", allow_blank=False); inw.add_data_validation(dv2); dv2.add('C4')
put(inw, 'B5', 'Bank name'); put(inw, 'C5', f"=IFERROR(INDEX('Peer Data'!$B$5:$B${LAST},MATCH($C$4,'Peer Data'!$A$5:$A${LAST},0)),\"Code not found\")", GREEN)
put(inw, 'B6', 'Period'); put(inw, 'C6', f"=IFERROR(INDEX('Peer Data'!$C$5:$C${LAST},MATCH($C$4,'Peer Data'!$A$5:$A${LAST},0)),\"\")", GREEN)
put(inw, 'B7', 'Months in income period'); put(inw, 'C7', 12, BLUE, fill=YEL); inw['C7'].comment = Comment('12 for full-year (Q4) figures; 3, 6 or 9 for interim quarters. Used to annualise profitability ratios.', 'Jamal Kanbarieh')
put(inw, 'B8', 'Sovereign ceiling'); put(inw, 'C8', 'None', BLUE, fill=YEL)
dv3 = DataValidation(type='list', formula1=f"=Assumptions!$E$19:$E${19 + len(SCALE)}", allow_blank=False); inw.add_data_validation(dv3); dv3.add('C8')
inw['C8'].comment = Comment('Optional: the final rating cannot be better than the sovereign ceiling. Syria has no public sovereign rating, so the default is None.', 'Jamal Kanbarieh')
hdr(inw, 10, ['Item', 'From Peer Data', 'Override', 'Value used', 'Unit'])
for j, (lab, _) in enumerate(FIELDS):
    r = 11 + j; col = L(4 + j)
    put(inw, f'A{r}', lab)
    put(inw, f'B{r}', f"=IFERROR(INDEX('Peer Data'!${col}$5:${col}${LAST},MATCH($C$4,'Peer Data'!$A$5:$A${LAST},0)),0)", GREEN, fmt=MN)
    put(inw, f'C{r}', None, BLUE, fmt=MN, fill=YEL)
    put(inw, f'D{r}', f'=IF(C{r}="",B{r},C{r})', BLACK, True, fmt=MN); put(inw, f'E{r}', 'SYP mn', '52514E')
for col, w in zip('ABCDE', [34, 18, 16, 18, 9]): inw.column_dimensions[col].width = w
IN = {lab: f'Inputs!$D${11 + j}' for j, (lab, _) in enumerate(FIELDS)}
TA, LQ, GL, NPL, ECL, DEP, CM, TL, EQ, RC, RWA, GOI, OPX, NI = [IN[f[0]] for f in FIELDS]
MON = 'Inputs!$C$7'

def ratio_formulas(g):  # g(x) -> reference for item x
    ta, lq, gl, npl, ecl, dep, cm, tl, eq, rc, rwa, goi, opx, ni = g
    return [
        f'=IF({rwa}<=0,"n.m.",{rc}/{rwa})',
        f'=IF({rwa}<=0,"n.m.",{eq}/{rwa})',
        f'=IF({gl}<=0,"n.m.",{npl}/{gl})',
        f'=IF(({eq}+{ecl})<=0,"n.m.",{npl}/({eq}+{ecl}))',
        f'=IF({rwa}<=0,"n.m.",{ni}*12/{MON}/{rwa})',
        f'=IF({rwa}<=0,"n.m.",({goi}-{opx})*12/{MON}/{rwa})',
        f'=IF({ta}<=0,"n.m.",({tl}-{dep}-{cm}-{lq})/{ta})',
        f'=IF({dep}<=0,"n.m.",{lq}/{dep})',
        f'=IF({goi}<=0,"n.m.",{opx}/{goi})',
    ]
def points_formula(x, a):
    A = f'Assumptions!$D${a}'; B = f'Assumptions!$E${a}'; C = f'Assumptions!$F${a}'; Dd = f'Assumptions!$G${a}'
    hi = f'IF({x}>={A},1,IF({x}>={B},2,IF({x}>={C},3,IF({x}>={Dd},4,5))))'
    lo = f'IF({x}<={A},1,IF({x}<={B},2,IF({x}<={C},3,IF({x}<={Dd},4,5))))'
    return f'=IF(ISNUMBER({x}),IF(Assumptions!$C${a}="Higher",{hi},{lo}),5)'

# ---------------- Scorecard ----------------
sc = sheet('Scorecard')
put(sc, 'A1', 'Credit scorecard', BLACK, True, sz=12)
put(sc, 'A2', '=Inputs!C5&"  ("&Inputs!C4&")  ·  "&Inputs!C6', GREEN, True, sz=11)
hdr(sc, 4, ['Ratio', 'Category', 'Value', 'Points (1–5)', 'Grade', 'Weight', 'Weighted points'])
rf = ratio_formulas([TA, LQ, GL, NPL, ECL, DEP, CM, TL, EQ, RC, RWA, GOI, OPX, NI])
for i in range(9):
    r = 5 + i; a = 5 + i
    put(sc, f'A{r}', f'=Assumptions!A{a}', GREEN); put(sc, f'B{r}', f'=Assumptions!B{a}', GREEN)
    put(sc, f'C{r}', rf[i], BLACK, fmt=PC2 if i in (4, 5) else PC, align='right')
    put(sc, f'D{r}', points_formula(f'C{r}', a), BLACK, align='center')
    put(sc, f'E{r}', f'=CHAR(64+D{r})', BLACK, True, align='center')
    put(sc, f'F{r}', f'=Assumptions!H{a}', GREEN, fmt=PC)
    put(sc, f'G{r}', f'=D{r}*F{r}', BLACK, fmt='0.000')
put(sc, 'F14', 'Score', BLACK, True); put(sc, 'G14', '=SUM(G5:G13)', BLACK, True, fmt='0.00')
for c in 'ABCDEFG': sc[f'{c}14'].border = Border(top=thin)
R = f'Assumptions!$B${SC0}:$B${SC1}'; SA = f'Assumptions!$A${SC0}:$A${SC1}'; PD = f'Assumptions!$C${SC0}:$C${SC1}'
res = [
 ('Weighted score (1 = strongest, 5 = weakest)', '=G14', '0.00'),
 ('Standalone rating', f'=INDEX({R},MATCH(G14,{SA},1))', None),
 ('Sovereign ceiling', '=Inputs!C8', None),
 ('Final rating', f'=IF(Inputs!C8="None",B18,INDEX({R},MAX(MATCH(B18,{R},0),MATCH(Inputs!C8,{R},0))))', None),
 ('Indicative 1-year PD', f'=INDEX({PD},MATCH(B20,{R},0))', PC2),
]
put(sc, 'A16', 'Result', BLACK, True, sz=11)
for i, (a, f_, fmt) in enumerate(res):
    r = 17 + i; put(sc, f'A{r}', a, BLACK, i == 3); c = put(sc, f'B{r}', f_, GREEN if i == 2 else BLACK, True, fmt=fmt, fill=OUT if i in (3, 4) else None)
put(sc, 'A24', 'Category view', BLACK, True, sz=11)
hdr(sc, 25, ['Category', 'Weight', 'Average points', 'Grade'])
for i, cat in enumerate(['Capital', 'Asset quality', 'Profitability', 'Liquidity', 'Efficiency']):
    r = 26 + i; put(sc, f'A{r}', cat)
    put(sc, f'B{r}', f'=SUMIF($B$5:$B$13,A{r},$F$5:$F$13)', BLACK, fmt=PC)
    put(sc, f'C{r}', f'=IFERROR(SUMIF($B$5:$B$13,A{r},$G$5:$G$13)/B{r},"")', BLACK, fmt='0.00')
    put(sc, f'D{r}', f'=IFERROR(CHAR(64+ROUND(C{r},0)),"")', BLACK, True, align='center')
for col, w in zip('ABCDEFG', [40, 15, 12, 13, 8, 9, 15]): sc.column_dimensions[col].width = w
sc.row_dimensions[4].height = 30

# ---------------- Peer Ranking ----------------
pr = sheet('Peer Ranking')
put(pr, 'A1', f'Peer ranking: every bank scored with the same model ({P})', BLACK, True, sz=12)
put(pr, 'A2', 'Ratios use "Peer Data" directly (overrides on "Inputs" do not apply here). Months in period are taken from Inputs!C7.', '52514E', i=True)
short = ['CAR', 'Equity/RWA', 'NPL ratio', 'NPL/(Eq+ECL)', 'NI/RWA', 'PPP/RWA', 'Mkt-fund gap', 'Liquid/deposits', 'CIR']
hdr(pr, 4, ['Code', 'Bank'] + short + [f'Pts {s}' for s in short] + ['Score', 'Rating', 'Rank'])
for i in range(NB):
    r = 5 + i; pdr = 5 + i
    put(pr, f'A{r}', f"='Peer Data'!A{pdr}", GREEN); put(pr, f'B{r}', f"='Peer Data'!B{pdr}", GREEN)
    g = [f"'Peer Data'!{L(4 + j)}{pdr}" for j in range(14)]
    for j, f_ in enumerate(ratio_formulas(g)):
        put(pr, f'{L(3 + j)}{r}', f_, BLACK, fmt=PC2 if j in (4, 5) else PC, align='right')
        put(pr, f'{L(12 + j)}{r}', points_formula(f'{L(3 + j)}{r}', 5 + j), BLACK, align='center')
    put(pr, f'U{r}', '=' + '+'.join(f'{L(12 + j)}{r}*Assumptions!$H${5 + j}' for j in range(9)), BLACK, True, fmt='0.00')
    put(pr, f'V{r}', f'=INDEX({R},MATCH(U{r},{SA},1))', BLACK, True, align='center')
    put(pr, f'W{r}', f'=RANK(U{r},$U$5:$U${LAST},1)', BLACK, align='center')
pr.column_dimensions['A'].width = 8; pr.column_dimensions['B'].width = 30
for j in range(3, 24): pr.column_dimensions[L(j)].width = 11
pr.row_dimensions[4].height = 42; pr.freeze_panes = 'C5'

wb._sheets = [wb[n] for n in ['Cover', 'Inputs', 'Scorecard', 'Peer Ranking', 'Assumptions', 'Peer Data']]; wb.active = 0
wb.properties.creator = 'Jamal Kanbarieh'; wb.properties.title = 'Bank Counterparty Credit Scoring Model'
wb.save('/home/claude/portfolio/projects/credit-scoring/Bank_Credit_Scoring_Model.xlsx')
print('saved', NB, 'banks')
